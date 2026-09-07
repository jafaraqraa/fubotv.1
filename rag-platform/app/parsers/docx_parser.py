import io
import docx
from app.parsers.base import BaseParser, ParsedDocument

class DOCXParser(BaseParser):
    def parse(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        doc = docx.Document(io.BytesIO(file_bytes))
        sections = []
        full_text = []
        current_heading = "General"
        current_section_lines = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if p.style.name.startswith("Heading"):
                if current_section_lines:
                    sections.append({
                        "title": current_heading,
                        "content": "\n".join(current_section_lines)
                    })
                    current_section_lines = []
                current_heading = text
            else:
                current_section_lines.append(text)
                full_text.append(text)

        if current_section_lines:
            sections.append({
                "title": current_heading,
                "content": "\n".join(current_section_lines)
            })

        doc_title = file_name.rsplit(".", 1)[0]
        return ParsedDocument(
            title=doc_title,
            content="\n".join(full_text),
            metadata={},
            sections=sections if sections else [{"title": "General", "content": "\n".join(full_text)}]
        )
